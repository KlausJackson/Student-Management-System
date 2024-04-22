import sqlite3
from PySide6.QtCore import (QCoreApplication, Signal, Qt)
from PySide6.QtWidgets import (QFileDialog, QMessageBox, QTableWidgetItem)
from PySide6.QtGui import (QFont)
from PySide6.QtWidgets import (QLineEdit, QMessageBox, QPushButton, QLabel, QDialog, 
                             QVBoxLayout, QComboBox, QGridLayout)
from Custom_Widgets.QCustomModals import QCustomModals
import src._icons_rc
import pandas as pd
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from docx import Document



class Func():
    def __init__(self, ui):
        super(Func, self).__init__()
        self.ui = ui
        self.database = None
        self.total = 0
        
        
    def load_data(self):
        database = self.database
        connection = sqlite3.connect(database)
        self.ui.table.setRowCount(0)
        result = connection.execute("SELECT * FROM student")
        self.total = connection.execute("SELECT COUNT(*) FROM student").fetchone()[0]
        self.ui.total.setText(QCoreApplication.translate("MainWindow", f"Total: {self.total}", None))
        
        for row_num, row_data in enumerate(result):
            self.ui.table.insertRow(row_num)
            for column_num, data in enumerate(row_data):
                self.ui.table.setItem(row_num, column_num, QTableWidgetItem(str(data)))
        
        for row in range(self.ui.table.rowCount()):
            for column in range(self.ui.table.columnCount()):
                item = self.ui.table.item(row, column)
                if item is not None:
                    content_font = QFont("Arial", 12) 
                    item.setFont(content_font)
        
        for column in range(self.ui.table.columnCount()):
            max_width = 0
            for row in range(self.ui.table.rowCount()):
                item = self.ui.table.item(row, column)
                if item is not None:
                    max_width = max(max_width, self.ui.table.fontMetrics().boundingRect(item.text()).width())
            self.ui.table.setColumnWidth(column, max_width + 70)
        connection.close()
        
        notify = QCustomModals.InformationModal(
            title = "Action Made",
            parent = self.ui.centralwidget,  
            position = 'top-right',
            closeIcon = u":/feather/icons/feather/x-circle.png", 
            description = "Database has been updated.", 
            isClosable = True, 
            duration = 3000 
        )   
        notify.show()
        
        
    def search(self):
        if self.database != None:
            self.dialog = FindDialog(self.database)  
            self.dialog.search_result.connect(self.show_search) 
            self.dialog.finished.connect(self.load_data) 
            self.dialog.show()  
        else:
            self.no_database()
            
    def show_search(self, rows):
        self.ui.table.setRowCount(0)
        for row_num, row_data in enumerate(rows):
            self.ui.table.insertRow(row_num)
            for column_num, data in enumerate(row_data):
                item = QTableWidgetItem(str(data))
                self.ui.table.setItem(row_num, column_num, item)
                
        for row in range(self.ui.table.rowCount()):
            for column in range(self.ui.table.columnCount()):
                item = self.ui.table.item(row, column)
                if item is not None:
                    content_font = QFont("Arial", 12) 
                    item.setFont(content_font)         
                         
        
    def no_database(self):
        notify = QCustomModals.WarningModal(
                    title="Warning!",
                    description="No database, please upload your database file first.",
                    parent = self.ui.centralwidget,  
                    position = 'top-right',
                    closeIcon = u":/feather/icons/feather/x-circle.png",  
                    isClosable = True, 
                    duration = 3000                
                )
        notify.show()   
        return            
        
    def add(self):
        if self.database != None:
            id, name = self.ui.id.text(), self.ui.name.text()
            if name == "" or id == "":
                notify = QCustomModals.WarningModal(
                    title="Warning!",
                    description="Name and ID can't be empty.",
                    parent = self.ui.centralwidget,  
                    position = 'top-right',
                    closeIcon = u":/feather/icons/feather/x-circle.png",  
                    isClosable = True, 
                    duration = 3000                
                )
                notify.show()   
                return
            
            course, mobile, note = self.ui.course.currentText(), self.ui.contact.text(), self.ui.note.text()
            connection = sqlite3.connect(self.database)    
            cursor = connection.cursor()
            cursor.execute("INSERT INTO student (id, name, course, mobile, note) VALUES (?, ?, ?, ?, ?)", 
                        (id, name, course, mobile, note))
            
            connection.commit()
            cursor.close()
            connection.close()
            self.load_data()
            self.ui.id.clear()
            self.ui.name.clear()
            self.ui.contact.clear()
            self.ui.note.clear()  
        else:
            self.no_database()
        
        
    def get_database(self):
        path, _ = QFileDialog.getOpenFileName(None, "Select Database file", "", "SQLite Database (*.db)")
        if path:
            self.database = path
            self.load_data()
        
    def edit(self):
        dialog = EditDialog(self.ui, self, self.database)
        dialog.exec()    
    def delete(self):
        dialog = DeleteDialog(self.ui, self, self.database)
        dialog.exec()   
        
        
    def get_data(self):
        if self.database != None:
            path = QFileDialog.getExistingDirectory(None, "Select where to save your file", "")
            if path:
                data = []
                for row in range(self.ui.table.rowCount()):
                    row_data = []
                    for column in range(self.ui.table.columnCount()):
                        item = self.ui.table.item(row, column)
                        row_data.append(item.text())
                    data.append(row_data)        
                return data, path
        else:
            self.no_database()
        
    def get_header(self):
        headers = []
        for column in range(self.ui.table.columnCount()):
            headers.append(self.ui.table.horizontalHeaderItem(column).text()) 
        return headers
        
    def pdf(self):
        result = self.get_data() 
        if result: 
            data, path = result
            data = [list(self.get_header())] + data
            file = SimpleDocTemplate(path + "/Students.pdf", pagesize=letter)
            table = Table(data)
            style = TableStyle([('GRID', (0,0), (-1,-1), 1, (0,0,0)), ('FONTSIZE', (0, 0), (-1, -1), 12)])             
            table.setStyle(style)  
            file.build([table])            
        
    def doc(self):
        result = self.get_data() 
        if result: 
            data, path = result
            file = Document()
            table = file.add_table(rows=len(data) + 1, cols=len(self.get_header()))
            for i, header in enumerate(self.get_header()):
                table.cell(0, i).text = header            
            for i, row in enumerate(data):
                for j, cell_text in enumerate(row):
                    table.cell(i, j).text = cell_text            
            file.save(path + "/Students.docx")
            
    def excel(self):
        result = self.get_data() 
        if result: 
            data, path = result
            df = pd.DataFrame(data)
            df.to_excel(path + "/Students.xlsx", index=False, header=self.get_header)      
                                  
    def csv(self):   
        result = self.get_data()
        if result:
            data, path = result
            df = pd.DataFrame(data)
            df.to_csv(path + "/Students.csv", index=False, header=self.get_header)    
                    

        
class EditDialog(QDialog):
    def __init__(self, ui, func, database):
        super().__init__()                    
        self.setWindowTitle("Edit Student Data")       
        self.setFixedHeight(250)
        self.setFixedWidth(370)
        self.ui, self.func, self.database = ui, func, database
       # self.setAttribute(Qt.WA_DeleteOnClose)

        
        i = self.ui.table.currentRow()
        name = self.ui.table.item(i, 1).text()            
        id = self.ui.table.item(i, 0).text()
        self.original_id = id
        contact = self.ui.table.item(i, 3).text()                        
        course = self.ui.table.item(i, 2).text()
        note = self.ui.table.item(i, 4).text()
                    
        layout = QVBoxLayout()
        self.id = QLineEdit(id)
        self.id.setPlaceholderText("ID")
        layout.addWidget(self.id)
        
        self.name = QLineEdit(name) 
        self.name.setPlaceholderText("Name")
        layout.addWidget(self.name)
        
        self.course_list = []
        for index in range(self.ui.course.count()):
            self.course_list.append(self.ui.course.itemText(index))     
        self.course = QComboBox() 
        self.course.addItems(self.course_list)
        self.course.setCurrentText(course)
        layout.addWidget(self.course)
        
        self.contact = QLineEdit(contact)
        self.contact.setPlaceholderText("Phone Number, Email, etc.")
        layout.addWidget(self.contact)
        
        self.note = QLineEdit(note)
        self.note.setPlaceholderText("Note")
        layout.addWidget(self.note)
        
        button = QPushButton("Update")        
        button.clicked.connect(self.update)
        layout.addWidget(button)
        self.setLayout(layout) 
        
    # def closeEvent(self, event):
    #     super().closeEvent(event)        
        
    def update(self):
        id, name, course, mobile = self.id.text(), self.name.text(), self.course.currentText(), self.contact.text()
        note = self.note.text()
        if name == "" or id == "":
            message = QMessageBox()
            message.setWindowTitle("ERROR")
            message.setText("Name or ID cannot be empty.")
            message.exec()
            return        
        
        connection = sqlite3.connect(self.database)                    
        cursor = connection.cursor()    
        cursor.execute("UPDATE student SET id = ?, name = ?, course = ?, mobile = ?, note = ? WHERE id = ?",
                       (id, name, course, mobile, note, self.original_id))    
        connection.commit()
        cursor.close()    
        connection.close()
        self.func.load_data() 
        
        notify = QCustomModals.InformationModal(
            title = "Action Made",
            parent = self.ui.centralwidget,  
            position = 'top-right',
            closeIcon = u":/feather/icons/feather/x-circle.png", 
            description = f"{name} has been updated.", 
            isClosable = True, 
            duration = 3000 
        )   
        notify.show()                 
        
        
        
class DeleteDialog(QDialog):
    def __init__(self, ui, func, database):
        super().__init__()                    
        self.setWindowTitle("Delete Student Data")      
        self.setFixedHeight(100)
        self.setFixedWidth(400)
        self.ui, self.func, self.database = ui, func, database                           
        self.setAttribute(Qt.WA_DeleteOnClose)           
        layout = QGridLayout()
        
        message = QLabel("Delete this student? This action cannot be reversed.")    
        y, n = QPushButton("Yes"), QPushButton("No")          
        layout.addWidget(message, 0, 0, 1, 2)
        layout.addWidget(y, 1, 0)
        layout.addWidget(n, 1, 1)   
        self.setLayout(layout)
        y.clicked.connect(self.delete)  
        n.clicked.connect(self.close)
        
    def closeEvent(self, event):
        super().closeEvent(event)           
                
    def delete(self):
        i = self.ui.table.currentRow()                
        name = self.ui.table.item(i, 1).text()            
        connection = sqlite3.connect(self.database)            
        cursor = connection.cursor()
        cursor.execute("DELETE FROM student WHERE name = ?", (name, ))
        connection.commit()
        cursor.close()
        connection.close()
        self.func.load_data()
        notify = QCustomModals.InformationModal(
            title = "Action Made",
            parent = self.ui.centralwidget,  
            position = 'top-right', 
            closeIcon = u":/feather/icons/feather/x-circle.png", 
            description = f"{name} has been delete from the database.", 
            isClosable = True, 
            duration = 3000 
        )
        notify.show()         
        self.close()
        
        

class FindDialog(QDialog):
    search_result = Signal(list)
    def __init__(self, database):
        super().__init__()
        self.setWindowTitle("Find Student")
        self.setFixedHeight(100)
        self.setFixedWidth(370)
        self.result = 0
        self.database = database
        
        layout = QVBoxLayout()
        self.find = QLineEdit()
        self.find.setPlaceholderText("Any information you want to look for.")
        self.find.textChanged.connect(self.search)
        layout.addWidget(self.find)
        
        self.report = QLabel(f"{self.result} matched.")
        layout.addWidget(self.report)        
        self.setLayout(layout)        

    def search(self):
        info = self.find.text()
        connection = sqlite3.connect(self.database)
        cursor = connection.cursor()
        rows = list(cursor.execute("SELECT * FROM student WHERE name LIKE ? OR id LIKE ? OR course LIKE ? OR mobile LIKE ? OR note LIKE ?"
                                   , ('%' + info + '%', '%' + info + '%', '%' + info + '%', '%' + info + '%', '%' + info + '%')))      
        self.result = len(rows)
        self.report.setText(f"{self.result} matched.")
        self.search_result.emit(rows)  
              
        
        
        