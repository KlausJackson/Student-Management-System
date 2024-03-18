from PyQt6.QtWidgets import (QApplication, QLabel, QPushButton, QMainWindow, QTableWidget,
                             QTableWidgetItem, QFileDialog, QToolBar, QStatusBar)
from PyQt6.QtGui import QIcon, QAction, QFont
import sys
import sqlite3
from actions import InsertDialog, DeleteDialog, FindDialog, EditDialog, AboutDialog
from docx import Document
import pandas as pd
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle



class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Student Management System")
        self.setWindowIcon(QIcon("icons/logo.png"))
        self.setGeometry(100, 100, 900, 600)
        font = QFont("Arial", 12)
        QApplication.instance().setFont(font)  
        self.total = 0 
        
        file = self.menuBar().addMenu("&File")
        doc, pdf = QAction("Save to Doc", self), QAction("Save to PDF", self)
        csv, excel = QAction("Save to CSV", self), QAction("Save to Excel", self)
        
        pdf.triggered.connect(self.to_pdf)
        doc.triggered.connect(self.to_doc)
        excel.triggered.connect(self.to_excel)
        csv.triggered.connect(self.to_csv)
        
        file.addAction(doc)
        file.addAction(excel)
        file.addAction(pdf)
        file.addAction(csv)
        
        edit = self.menuBar().addMenu("&Edit")
        add = QAction(QIcon("icons/add.png"), "Add", self)
        find = QAction(QIcon("icons/search.png"), "Find", self)
        find.triggered.connect(self.find)
        add.triggered.connect(self.insert)
        edit.addAction(add)         
        edit.addAction(find)
        
        help = self.menuBar().addMenu("&Help")
        about = QAction("About", self)
        about.triggered.connect(self.about_info)
        help.addAction(about)
        
        self.table = QTableWidget()
        self.table.setColumnCount(5)
        self.headers = ("ID", "Name", "Course", "Contact Info", "Note")
        self.table.setHorizontalHeaderLabels(self.headers)
        self.setCentralWidget(self.table)
        self.load_data()
        
        toolbar = QToolBar()
        toolbar.setMovable(True)
        self.addToolBar(toolbar)
        toolbar.addAction(add)
        toolbar.addAction(find)
        
        self.statusbar = QStatusBar()
        self.setStatusBar(self.statusbar)
        total_student = QLabel(f"Total: {self.total}")
        self.statusbar.addWidget(total_student)
        self.table.cellClicked.connect(self.cell_clicked)
       
    def cell_clicked(self):   
        edit_button = QPushButton("Edit")  
        delete_button = QPushButton("Delete")
        delete_button.clicked.connect(self.delete) 
        edit_button.clicked.connect(self.edit_info)
        
        kids = self.findChildren(QPushButton)
        if kids:
            for kid in kids:
                self.statusbar.removeWidget(kid)
        self.statusbar.addWidget(edit_button)
        self.statusbar.addWidget(delete_button)
         

    def edit_info(self):
        dialog = EditDialog(sms)
        dialog.exec()    
    def delete(self):
        dialog = DeleteDialog(sms)
        dialog.exec()    
    def about_info(self):
        dialog = AboutDialog()
        dialog.exec()
        
    def load_data(self):
        connection = sqlite3.connect("E:\\Work\\Python\\projects\\Student Management System\\database.db")
        self.table.setRowCount(0)
        result = connection.execute("SELECT * FROM student")
        self.total = connection.execute("SELECT COUNT(*) FROM student").fetchone()[0]
        
        for row_num, row_data in enumerate(result):
            self.table.insertRow(row_num)
            for column_num, data in enumerate(row_data):
                self.table.setItem(row_num, column_num, QTableWidgetItem(str(data)))

        for column in range(self.table.columnCount()):
            max_width = 0
            for row in range(self.table.rowCount()):
                item = self.table.item(row, column)
                if item is not None:
                    max_width = max(max_width, self.table.fontMetrics().boundingRect(item.text()).width())
            self.table.setColumnWidth(column, max_width + 70)

        connection.close()
        
    def insert(self):
        dialog = InsertDialog(sms)    
        dialog.exec()
    def find(self):
        self.dialog = FindDialog()  
        self.dialog.search_result.connect(self.show_search) 
        self.dialog.finished.connect(self.load_data) 
        self.dialog.show()
        
    def show_search(self, rows):
        self.table.setRowCount(0)
        for row_num, row_data in enumerate(rows):
            self.table.insertRow(row_num)
            for column_num, data in enumerate(row_data):
                item = QTableWidgetItem(str(data))
                self.table.setItem(row_num, column_num, item)             
        
    def get_data(self):
        path = QFileDialog.getExistingDirectory(None, "Select where to save your file", "")
        if path:
            data = []
            for row in range(self.table.rowCount()):
                row_data = []
                for column in range(self.table.columnCount()):
                    item = self.table.item(row, column)
                    row_data.append(item.text())
                data.append(row_data)        
            return data, path
        
    def to_pdf(self):
        result = self.get_data() 
        if result: 
            data, path = result
            data = [list(self.headers)] + data
            file = SimpleDocTemplate(path + "/Students.pdf", pagesize=letter)
            table = Table(data)
            style = TableStyle([('GRID', (0,0), (-1,-1), 1, (0,0,0)), ('FONTSIZE', (0, 0), (-1, -1), 12)])             
            table.setStyle(style)  
            file.build([table])            
        
    def to_doc(self):
        result = self.get_data() 
        if result: 
            data, path = result
            file = Document()
            table = file.add_table(rows=len(data) + 1, cols=len(self.headers))
            for i, header in enumerate(self.headers):
                table.cell(0, i).text = header            
            for i, row in enumerate(data):
                for j, cell_text in enumerate(row):
                    table.cell(i, j).text = cell_text            
            file.save(path + "/Students.docx")
            
    def to_excel(self):
        result = self.get_data() 
        if result: 
            data, path = result
            df = pd.DataFrame(data)
            df.to_excel(path + "/Students.xlsx", index=False, header=self.headers)      
                                   
    def to_csv(self):
        result = self.get_data()
        if result:
            data, path = result
            df = pd.DataFrame(data)
            df.to_csv(path + "/Students.xlsx", index=False, header=self.headers)    
                               
            
        
app = QApplication(sys.argv)
sms = MainWindow()
sms.show()
sys.exit(app.exec())







